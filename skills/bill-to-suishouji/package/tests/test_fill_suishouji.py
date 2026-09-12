import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from scripts.fill_suishouji import (
    build_suishouji_entry,
    classify_expense,
    classify_income,
    fill_sheet,
    fill_suishouji_template,
    filter_rows_from_start,
    manual_supplement_reminders,
    map_account,
    write_report,
)


class FillSuishoujiTest(unittest.TestCase):
    def test_filter_rows_from_start_includes_exact_timestamp(self):
        rows = [
            {"交易时间": "2026-06-02 08:41:28"},
            {"交易时间": "2026-06-02 08:41:29"},
            {"交易时间": "2026-06-02 08:41:30"},
        ]

        result = filter_rows_from_start(rows, "2026-06-02 08:41:29")

        self.assertEqual([row["交易时间"] for row in result], ["2026-06-02 08:41:29", "2026-06-02 08:41:30"])

    def test_map_account_uses_taxonomy_credit_card_cutoff(self):
        before_cutoff = {"交易时间": "2026-06-14 23:59:59", "收/付款方式": "招商银行信用卡(6020)", "收/支": "支出"}
        on_cutoff = {"交易时间": "2026-06-15 00:00:00", "收/付款方式": "招商银行信用卡(6020)", "收/支": "支出"}
        late = {"交易时间": "2026-06-16 09:00:00", "收/付款方式": "招商银行信用卡(6020)&国家补贴"}
        debit = {"交易时间": "2026-06-16 09:00:00", "收/付款方式": "招商银行储蓄卡(0329)", "收/支": "支出"}

        self.assertEqual(map_account(before_cutoff), "招行7.2还款")
        self.assertEqual(map_account(on_cutoff), "招行8.2还款")
        self.assertEqual(map_account(late), "招行8.2还款")
        self.assertEqual(map_account(debit), "招行")

    def test_classifies_common_expense_and_income_examples(self):
        self.assertEqual(classify_expense({"备注": "luckin coffee | 订单付款", "交易对象": "luckin coffee"}), ("食品酒水", "☕️咖啡"))
        self.assertEqual(classify_expense({"备注": "菜鸟裹裹商家寄件 | 寄件费", "交易对象": "菜鸟裹裹商家寄件"}), ("生活不易", "邮寄费"))
        self.assertEqual(classify_income({"备注": "[航延航意险（大家计划二）]理赔", "交易对象": "[航延航意险（大家计划二）]理赔", "金额": 60}), ("其他收入", "保险"))

    def test_classifies_obvious_review_cases(self):
        self.assertEqual(classify_expense({"备注": "wildlimes咖啡 | wildlimes", "交易对象": "wildlimes咖啡"}), ("食品酒水", "☕️咖啡"))
        self.assertEqual(classify_expense({"备注": "奈雪的茶 | 珠海机场PRO店", "交易对象": "奈雪的茶"}), ("食品酒水", "🥤奶茶果茶"))
        self.assertEqual(classify_expense({"备注": "全棉时代一次性内裤男女旅行出差", "交易对象": "天**"}), ("买买买", "衣服裤子"))

    def test_classifies_202607_manual_review_rules(self):
        cases = [
            ({"交易对象": "Floury", "商品说明": "Floury", "备注": "Floury | Floury"}, ("食品酒水", "早餐")),
            ({"交易对象": "美团平台商户", "商品说明": "章鱼陶陶·陶艺银饰·拼豆调香·手工DIY", "备注": "章鱼陶陶·陶艺银饰·拼豆调香·手工DIY"}, ("玩得开心", "其他活动")),
            ({"交易对象": "薛记炒货", "商品说明": "J169-消费", "备注": "薛记炒货 | J169-消费"}, ("买买买", "家庭购物")),
            ({"交易对象": "沃尔玛（中国）投资有限公司", "商品说明": "山姆自助购", "备注": "沃尔玛（中国）投资有限公司 | 山姆自助购"}, ("买买买", "家庭购物")),
            ({"交易对象": "tb**7", "商品说明": "Planwith Sign三角多色笔", "备注": "Planwith Sign三角多色笔"}, ("学得开心", "文具装备")),
            ({"交易对象": "盒马", "商品说明": "盒马 原味酸奶 1.5kg等多件", "备注": "盒马 原味酸奶 1.5kg等多件"}, ("买买买", "家庭购物")),
            ({"交易对象": "Yuansfer Inc", "商品说明": "App Store & iTunes US - Pockyt GiftCard", "备注": "Yuansfer Inc | App Store & iTunes US - Pockyt GiftCard"}, ("生活不易", "会员费")),
            ({"交易对象": "珠海市高新区畔秋山手作店", "商品说明": "畔秋山手作店-消费", "备注": "畔秋山手作店-消费"}, ("食品酒水", "饭堂")),
            ({"交易对象": "中国南方航空", "商品说明": "B2607080003190", "备注": "中国南方航空 | B2607080003190"}, ("玩得开心", "旅游度假")),
            ({"交易对象": "浙江空港精选经营有限公司", "商品说明": "疯狂动物城-萌友集结系列毛绒手机绳盲盒3", "备注": "疯狂动物城-萌友集结系列毛绒手机绳盲盒3"}, ("玩得开心", "礼物")),
            ({"交易对象": "CEOL x OFF SPACE", "商品说明": "Ceol Coffee", "备注": "CEOL x OFF SPACE | Ceol Coffee"}, ("食品酒水", "☕️咖啡")),
        ]

        for row, expected in cases:
            with self.subTest(row=row):
                self.assertEqual(classify_expense(row), expected)

    def test_classifies_20260803_manual_review_rules(self):
        cases = [
            ({"备注": "甜津美 | 总共消费:52.0"}, ("买买买", "家庭购物")),
            ({"备注": "茉莉奶白 | 【打包】_珠海香洲玖洲道购物公园店"}, ("食品酒水", "🥤奶茶果茶")),
            ({"备注": "沾叔港式小吃 | 沾叔港式小吃"}, ("买买买", "家庭购物")),
            ({"备注": "名创优品 | 名创优品广东珠海中海富华里店-消费"}, ("买买买", "家庭购物")),
            ({"备注": "无印良品MUJI | 无印良品MUJI-珠海优特汇-POS1"}, ("买买买", "护肤品")),
            ({"备注": "NITORI | 宜得利珠海优特广场店"}, ("买买买", "家庭购物")),
            ({"备注": "盒马 | 盒马 奶盐味苏打饼干 680g等多件"}, ("买买买", "家庭购物")),
            ({"备注": "7ELEVEn | 多点订单：994010200236929"}, ("食品酒水", "早餐")),
            ({"备注": "丰 e 足食智慧零售 | 先购后付"}, ("食品酒水", "水果零食")),
            ({"备注": "东风日产 | 维保工单支付"}, ("行车交通", "私家车费用")),
            ({"备注": "TIMHORTONS | 天好深圳咖啡餐饮有限公司"}, ("食品酒水", "☕️咖啡")),
            ({"备注": "美洲城(澳門)戶外用品一人有限公司 | shopping"}, ("买买买", "衣服裤子")),
            ({"备注": "东风日产平台商户 | 商品"}, ("行车交通", "私家车费用")),
            ({"备注": "盒马 | 绣球 冰川蓝 1支等3类商品"}, ("买买买", "家庭购物")),
            ({"备注": "小红书 | 小红书订单：P800353460897401331"}, ("食品酒水", "外出吃饭")),
        ]

        for row, expected in cases:
            with self.subTest(row=row):
                self.assertEqual(classify_expense(row), expected)

        control_cases = [
            ({"备注": "7ELEVEn | 便利店消费"}, ("食品酒水", "水果零食")),
            ({"备注": "小红书 | 手账贴纸订单"}, ("买买买", "其他")),
            ({"备注": "盒马 | 【鲜果切】麒麟西瓜"}, ("食品酒水", "水果零食")),
        ]
        for row, expected in control_cases:
            with self.subTest(control=row):
                self.assertEqual(classify_expense(row), expected)

    def test_build_entry_matches_suishouji_columns(self):
        row = {
            "交易时间": "2026-06-02 08:41:29",
            "交易对象": "luckin coffee",
            "备注": "luckin coffee | 订单付款",
            "收/支": "支出",
            "金额": 10.9,
            "收/付款方式": "零钱",
        }

        entry = build_suishouji_entry(row)

        self.assertEqual(entry, ["支出", "2026-06-02 08:41:29", "食品酒水", "☕️咖啡", "微信钱", 10.9, "", "", "", "luckin coffee | 订单付款"])

    def test_builds_nasdaq_purchase_as_transfer(self):
        row = {
            "交易时间": "2026-08-21 10:23:25",
            "收/支": "转账",
            "金额": 5,
            "收/付款方式": "交通银行储蓄卡(8984)",
            "备注": "蚂蚁财富 | 广发纳斯达克100ETF联接(QDII)A-买入",
        }

        entry = build_suishouji_entry(row)

        self.assertEqual(entry, ["转账", "2026-08-21 10:23:25", "交行卡", "支付宝纳指", 5.0, "", "", "", "蚂蚁财富 | 广发纳斯达克100ETF联接(QDII)A-买入"])

    def test_builds_gold_purchase_as_transfer(self):
        row = {
            "交易时间": "2026-08-19 09:10:08",
            "收/支": "转账",
            "金额": 100,
            "收/付款方式": "",
            "备注": "南方上海金ETF联接C | 南方上海金ETF联接C-买入",
        }

        entry = build_suishouji_entry(row)

        self.assertEqual(entry, ["转账", "2026-08-19 09:10:08", "微信钱", "支付宝理财-黄金", 100.0, "", "", "", "南方上海金ETF联接C | 南方上海金ETF联接C-买入"])

    def test_fill_sheet_replaces_example_row_with_real_data(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["交易类型", "日期", "一级分类", "二级分类", "账户", "金额", "成员", "商家", "项目", "备注"])
        sheet.append(["支出", "2012-09-14 03:00:09", "休闲娱乐", "运动健身", "银行卡", 3700, "", "", "", "示例"])
        sheet.append(["", "", "", "", "", "", "", "", "", ""])

        fill_sheet(
            sheet,
            [["支出", "2026-06-02 08:41:29", "食品酒水", "☕️咖啡", "微信钱", 10.9, "", "", "", "luckin coffee | 订单付款"]],
        )

        self.assertEqual(sheet.cell(2, 2).value, "2026-06-02 08:41:29")
        self.assertNotEqual(sheet.cell(2, 2).value, "2012-09-14 03:00:09")
        self.assertEqual(sheet.cell(3, 1).value, None)

    def test_fill_sheet_clears_example_row_when_no_entries(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["交易类型", "日期", "转出账户", "转入账户", "金额", "成员", "商家", "项目", "备注"])
        sheet.append(["转账", "2012-09-14 03:00:09", "微信钱包", "工行309", 3700, "", "", "", "示例"])

        fill_sheet(sheet, [])

        self.assertEqual([sheet.cell(2, col).value for col in range(1, 10)], [None] * 9)

    def test_template_fill_clears_transfer_example_sheet(self):
        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            bills_path = tmp_path / "bills.xlsx"
            template_path = tmp_path / "template.xlsx"
            output_path = tmp_path / "output.xlsx"
            report_path = tmp_path / "report.md"

            bills = Workbook()
            bills_sheet = bills.active
            bills_sheet.title = "合并账单"
            bills_sheet.append(["交易时间", "交易对象", "商品说明", "收/支", "金额", "收/付款方式", "备注"])
            bills_sheet.append(["2026-06-02 08:41:29", "luckin coffee", "订单付款", "支出", 10.9, "零钱", "luckin coffee | 订单付款"])
            bills.save(bills_path)

            template = Workbook()
            expense = template.active
            expense.title = "支出"
            expense.append(["交易类型", "日期", "一级分类", "二级分类", "支出账户", "金额", "成员", "商家", "项目", "备注"])
            expense.append(["支出", "2012-09-14 03:00:09", "休闲娱乐", "运动健身", "银行卡", 3700, "", "", "", "示例"])
            income = template.create_sheet("收入")
            income.append(["交易类型", "日期", "一级分类", "二级分类", "收入账户", "金额", "成员", "商家", "项目", "备注"])
            income.append(["收入", "2012-09-14 03:00:09", "休闲娱乐", "运动健身", "银行卡", 3700, "", "", "", "示例"])
            transfer = template.create_sheet("转账")
            transfer.append(["交易类型", "日期", "转出账户", "转入账户", "金额", "成员", "商家", "项目", "备注"])
            transfer.append(["转账", "2012-09-14 03:00:09", "微信钱包", "工行309", 3700, "", "", "", "示例"])
            template.save(template_path)

            fill_suishouji_template(bills_path, template_path, output_path, report_path, "2026-06-02 08:41:29")

            result = load_workbook(output_path, read_only=True, data_only=True)
            self.assertEqual(result["支出"].cell(2, 2).value, "2026-06-02 08:41:29")
            self.assertEqual([result["收入"].cell(2, col).value for col in range(1, 11)], [None] * 10)
            self.assertEqual([result["转账"].cell(2, col).value for col in range(1, 10)], [None] * 9)

    def test_report_reminds_account_check_before_manual_review(self):
        with TemporaryDirectory() as tmp:
            report_path = Path(tmp) / "report.md"
            rows = [
                {
                    "交易时间": "2026-06-02 08:41:29",
                    "交易对象": "luckin coffee",
                    "商品说明": "订单付款",
                    "备注": "luckin coffee | 订单付款",
                    "收/支": "支出",
                    "金额": 10.9,
                    "收/付款方式": "零钱",
                }
            ]

            write_report(rows, Path("output.xls"), report_path, "2026-06-02 08:41:29")

            report = report_path.read_text(encoding="utf-8")
            self.assertIn("先检查账户是否准确", report)
            self.assertIn("账户确认后再进行手工分类校对", report)

    def test_manual_supplement_reminders_follow_bill_date_range(self):
        rows = [{"交易时间": "2026-07-26 12:00:00"}]

        reminders = manual_supplement_reminders(rows, "2026-07-04 00:00:00")

        self.assertIn("2026-07-05 前后：确认工资收入是否已手动补录。", reminders)
        self.assertIn("2026-07-13 左右：确认公积金收入是否已手动补录。", reminders)
        self.assertIn("2026-07-05 左右：确认房租支出是否已手动补录。", reminders)
        self.assertIn("2026-07-25 左右：确认房租支出是否已手动补录。", reminders)
        self.assertNotIn("2026-08-05 前后：确认工资收入是否已手动补录。", reminders)

    def test_report_includes_manual_supplement_reminders(self):
        with TemporaryDirectory() as tmp:
            report_path = Path(tmp) / "report.md"
            rows = [
                {
                    "交易时间": "2026-07-26 12:00:00",
                    "交易对象": "luckin coffee",
                    "商品说明": "订单付款",
                    "备注": "luckin coffee | 订单付款",
                    "收/支": "支出",
                    "金额": 10.9,
                    "收/付款方式": "零钱",
                }
            ]

            write_report(rows, Path("output.xls"), report_path, "2026-07-04 00:00:00")

            report = report_path.read_text(encoding="utf-8")
            self.assertIn("## 手动补记提醒", report)
            self.assertIn("确认工资收入是否已手动补录", report)
            self.assertIn("确认公积金收入是否已手动补录", report)
            self.assertIn("确认房租支出是否已手动补录", report)


if __name__ == "__main__":
    unittest.main()
