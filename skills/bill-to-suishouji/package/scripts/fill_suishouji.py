from __future__ import annotations

import argparse
import calendar
import os
import shutil
import subprocess
import tempfile
from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


EXPENSE_CATEGORY_PAIRS = {
    ("食品酒水", "☕️咖啡"),
    ("食品酒水", "🥤奶茶果茶"),
    ("食品酒水", "早餐"),
    ("食品酒水", "饭堂"),
    ("食品酒水", "外出吃饭"),
    ("食品酒水", "🧃柠檬茶可乐"),
    ("食品酒水", "水果零食"),
    ("食品酒水", "🥛酸奶牛奶凉茶健康的东西"),
    ("食品酒水", "在家🍳做饭"),
    ("食品酒水", "请客"),
    ("玩得开心", "谈恋爱"),
    ("玩得开心", "旅游度假"),
    ("玩得开心", "礼物"),
    ("玩得开心", "电影"),
    ("玩得开心", "玩具"),
    ("玩得开心", "回家机票"),
    ("玩得开心", "其他活动"),
    ("买买买", "家庭购物"),
    ("买买买", "衣服裤子"),
    ("买买买", "化妆品"),
    ("买买买", "护肤品"),
    ("买买买", "花💐"),
    ("买买买", "代购"),
    ("买买买", "🏠格力海岸"),
    ("买买买", "快乐的东西（2020算在玩得开心）"),
    ("买买买", "其他"),
    ("行车交通", "私家车费用"),
    ("行车交通", "公共交通"),
    ("行车交通", "打车租车"),
    ("学得开心", "心理咨询"),
    ("学得开心", "文具装备"),
    ("学得开心", "书报杂志"),
    ("学得开心", "培训进修"),
    ("学得开心", "自习费用"),
    ("学得开心", "惩罚"),
    ("生活不易", "会员费"),
    ("生活不易", "健身房"),
    ("生活不易", "房租"),
    ("生活不易", "手机费"),
    ("生活不易", "看病吃药"),
    ("生活不易", "邮寄费"),
    ("生活不易", "礼金"),
    ("其他杂项", "其他支出"),
}

INCOME_CATEGORY_PAIRS = {
    ("智力活动", "副业收入"),
    ("智力活动", "人才补贴"),
    ("智力活动", "奖金收入"),
    ("智力活动", "工资收入"),
    ("智力活动", "红包"),
    ("智力活动", "公积金"),
    ("其他收入", "Aa制还钱"),
    ("其他收入", "理财收入"),
    ("其他收入", "礼金收入"),
    ("其他收入", "中奖收入"),
    ("其他收入", "保险"),
    ("其他收入", "🏠房租"),
    ("其他收入", "他人借钱买🏠"),
    ("其他收入", "其他收入"),
}

VALID_ACCOUNTS = {
    "交行卡",
    "招行7.2还款",
    "招行8.2还款",
    "招行",
    "微信钱",
    "交行6.8还",
    "交行 5.8日还",
    "余额宝",
    "民生银行",
    "蚂蚁12.8日还",
    "现金1",
    "民生银行-9.8还",
    "支付宝",
    "支付宝纳指",
    "支付宝理财-黄金",
    "华润银行",
    "中国银行36171",
    "华润银行9390",
    "蚂蚁1.8日还",
    "交通银行-理财账户",
    "源点",
    "去哪儿☕️",
    "星巴克☕️",
    "富民（金山）",
}

def resolve_soffice() -> str:
    configured = os.environ.get("SUISHOUJI_SOFFICE")
    if configured:
        return configured
    found = shutil.which("soffice")
    if found:
        return found
    raise RuntimeError(
        "LibreOffice soffice not found. Install it or set SUISHOUJI_SOFFICE."
    )
ACCOUNT_REVIEW_REMINDER = "下一步请先检查账户是否准确，账户确认后再进行手工分类校对。"
GOLD_INVESTMENT_ACCOUNT = "支付宝理财-黄金"
NASDAQ_INVESTMENT_ACCOUNT = "支付宝纳指"
MANUAL_SUPPLEMENT_WINDOW_DAYS = 2
MANUAL_SUPPLEMENT_ITEMS = [
    (5, "前后", "确认工资收入是否已手动补录。"),
    (13, "左右", "确认公积金收入是否已手动补录。"),
    (5, "左右", "确认房租支出是否已手动补录。"),
    (25, "左右", "确认房租支出是否已手动补录。"),
]


def parse_dt(value: str | datetime) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d %H:%M:%S")


def filter_rows_from_start(rows: list[dict], start_time: str) -> list[dict]:
    start = parse_dt(start_time)
    return [row for row in rows if parse_dt(row["交易时间"]) >= start]


def month_starts_between(start: date, end: date) -> list[date]:
    current = date(start.year, start.month, 1)
    last = date(end.year, end.month, 1)
    months = []
    while current <= last:
        months.append(current)
        if current.month == 12:
            current = date(current.year + 1, 1, 1)
        else:
            current = date(current.year, current.month + 1, 1)
    return months


def manual_supplement_reminders(rows: list[dict], start_time: str) -> list[str]:
    start = parse_dt(start_time).date()
    row_dates = [parse_dt(row["交易时间"]).date() for row in rows if row.get("交易时间")]
    end = max(row_dates) if row_dates else start
    if end < start:
        end = start

    reminders = []
    for month_start in month_starts_between(start, end):
        _, last_day = calendar.monthrange(month_start.year, month_start.month)
        for day, label, message in MANUAL_SUPPLEMENT_ITEMS:
            target = date(month_start.year, month_start.month, min(day, last_day))
            window_start = target - timedelta(days=MANUAL_SUPPLEMENT_WINDOW_DAYS)
            window_end = target + timedelta(days=MANUAL_SUPPLEMENT_WINDOW_DAYS)
            if window_start <= end and window_end >= start:
                reminders.append(f"{target:%Y-%m-%d} {label}：{message}")
    return reminders


def normalized_text(row: dict) -> str:
    return " ".join(
        str(row.get(key) or "")
        for key in ["交易对象", "商品说明", "交易类型", "备注", "原始备注"]
    ).lower()


def has_any(text: str, keywords: list[str]) -> bool:
    return any(keyword.lower() in text for keyword in keywords)


def category(pair: tuple[str, str]) -> tuple[str, str]:
    if pair not in EXPENSE_CATEGORY_PAIRS and pair not in INCOME_CATEGORY_PAIRS:
        raise ValueError(f"分类不存在于分类体系: {pair[0]} / {pair[1]}")
    return pair


def classify_expense(row: dict) -> tuple[str, str]:
    text = normalized_text(row)
    counterparty = str(row.get("交易对象") or "").lower()

    if has_any(counterparty, ["luckin", "manner", "星巴克", "瑞幸", "peetscoffee", "peet", "咖啡", "ceol x off space", "timhortons"]) or has_any(text, ["ceol coffee", "timhortons"]):
        return category(("食品酒水", "☕️咖啡"))
    if has_any(text, ["喜茶", "1点点", "古茗", "comebuy", "甘杯", "胚子奶鲜", "诚记糖水", "玲姐甜品", "born與生", "古春堂", "茶语", "茶理", "茶月山", "tamkoko", "泰柯茶园", "奈雪", "冬瓜茶", "三号椰", "壶见", "茉莉奶白"]):
        return category(("食品酒水", "🥤奶茶果茶"))
    if has_any(text, ["金膳餐饮", "饭堂", "食堂", "畔秋山手作店"]):
        return category(("食品酒水", "饭堂"))
    if has_any(text, ["早餐", "早点", "肠粉", "小笼包", "floury"]) or (("7eleven" in text or "7-eleven" in text) and "多点订单" in text):
        return category(("食品酒水", "早餐"))
    if has_any(text, ["柠檬茶", "王老吉", "可乐", "雪碧"]):
        return category(("食品酒水", "🧃柠檬茶可乐"))
    if has_any(text, ["薛记炒货", "沃尔玛（中国）投资有限公司", "山姆自助购", "盒马 原味酸奶", "盒马 奶盐味苏打饼干", "绣球 冰川蓝", "甜津美", "沾叔港式小吃", "名创优品", "nitori", "宜得利"]):
        return category(("买买买", "家庭购物"))
    if has_any(text, ["盒马", "山姆", "钱大妈", "好一棵树", "湖畔市集", "7-eleven", "7eleven", "美宜佳", "赛壹便利店", "丰e足食", "丰 e 足食", "丰宜", "广东乐驿", "来福饼家", "鲍师傅", "凤梨酥", "多点订单", "便利"]):
        return category(("食品酒水", "水果零食"))
    if has_any(text, ["话梅油柑", "薄荷茶", "酸奶", "鲜奶", "凉茶"]):
        return category(("食品酒水", "🥛酸奶牛奶凉茶健康的东西"))
    if has_any(text, ["散养草鸡蛋", "菜市场", "生鲜", "鸡蛋"]):
        return category(("食品酒水", "在家🍳做饭"))
    if has_any(text, ["请客"]):
        return category(("食品酒水", "请客"))

    if has_any(text, ["中国石油", "中油碧辟", "中石化", "加油站", "停车", "捷停车", "顺易通", "速停车", "etc", "华发物业停车", "粤海金融中心", "通行费", "东风日产", "维保工单"]):
        return category(("行车交通", "私家车费用"))
    if has_any(text, ["地铁", "公交", "高铁", "火车", "12306", "滴滴", "嘀嘀", "出租车", "曹操出行", "t3出行", "哈啰出行"]):
        return category(("行车交通", "公共交通"))
    if has_any(text, ["打车", "租车"]):
        return category(("行车交通", "打车租车"))

    if has_any(text, ["内裤", "全棉时代"]):
        return category(("买买买", "衣服裤子"))

    if has_any(text, ["飞猪", "签证", "旅行", "旅行险", "携程", "机票", "酒店", "民宿", "门票", "景区", "度假", "清远长隆", "中国南方航空", "南方航空"]):
        return category(("玩得开心", "旅游度假"))
    if has_any(text, ["礼物", "相机盲盒", "泡泡玛特-礼物", "浙江空港精选经营有限公司", "疯狂动物城", "毛绒手机绳盲盒"]):
        return category(("玩得开心", "礼物"))
    if has_any(text, ["电影", "影院", "imax", "万达影城", "cgv"]):
        return category(("玩得开心", "电影"))
    if has_any(text, ["泡泡玛特", "pop mart"]):
        return category(("玩得开心", "玩具"))
    if has_any(text, ["回家机票"]):
        return category(("玩得开心", "回家机票"))
    if has_any(text, ["章鱼陶陶", "陶艺银饰", "拼豆调香", "手工diy"]):
        return category(("玩得开心", "其他活动"))
    if has_any(text, ["谈恋爱"]):
        return category(("玩得开心", "谈恋爱"))

    if has_any(text, ["童言心理", "心理咨询", "咨询"]):
        return category(("学得开心", "心理咨询"))
    if has_any(text, ["阿里云", "百炼", "kimi", "扣子", "印丫丫", "planwith sign", "三角多色笔", "中性笔"]):
        return category(("学得开心", "文具装备"))
    if has_any(text, ["得到", "微信读书", "知乎", "知识星球"]):
        return category(("学得开心", "书报杂志"))
    if has_any(text, ["培训", "进修"]):
        return category(("学得开心", "培训进修"))
    if has_any(text, ["自习"]):
        return category(("学得开心", "自习费用"))
    if has_any(text, ["惩罚"]):
        return category(("学得开心", "惩罚"))

    if has_any(text, ["菜鸟裹裹", "顺丰速运", "快递", "邮寄", "寄件费"]):
        return category(("生活不易", "邮寄费"))
    if has_any(text, ["apple.com", "会员", "wps", "滴答清单", "yuansfer", "pockyt giftcard", "ai会员", "gpt", "理发", "剪头发", "晨雨发型"]):
        return category(("生活不易", "会员费"))
    if has_any(text, ["正圆城运", "健身房", "健身"]):
        return category(("生活不易", "健身房"))
    if has_any(text, ["房租", "物业管理费", "物业费"]):
        return category(("生活不易", "房租"))
    if has_any(text, ["手机费", "话费", "中国移动", "中国联通", "中国电信"]):
        return category(("生活不易", "手机费"))
    if has_any(text, ["看病", "医院", "药店", "买药"]):
        return category(("生活不易", "看病吃药"))
    if has_any(text, ["份子钱", "礼金", "婚丧"]):
        return category(("生活不易", "礼金"))

    if has_any(text, ["肯德基", "kfc", "麦当劳", "汉堡王", "必胜客", "一番出汁", "寿喜烧", "大墩荣记", "天记茶餐厅", "茶餐厅", "黑手制面", "naive理想国", "顺德", "三叔粥铺", "月畔轩", "太二", "海底捞", "九毛九", "点都德", "广州酒家", "炳胜", "陶陶居", "南京大牌档", "外婆家", "绿茶餐厅", "西贝", "呷哺呷哺", "凑凑", "怂火锅", "柴火鸡", "美团外卖", "大众点评", "饿了么", "外卖订单", "烧烤", "烤鱼", "拉面", "渔记餐室", "寿司郎", "炭一点", "苏晓玥", "缓归", "兜记", "大福星鱼丸", "汐见", "小红书订单"]):
        return category(("食品酒水", "外出吃饭"))

    if has_any(text, ["家庭购物", "家居", "家电", "日用品", "洗衣液", "纸巾", "牙膏", "牙刷", "洗发水", "科沃斯", "扫地机器人", "抹布", "薛记炒货", "沃尔玛（中国）投资有限公司", "山姆自助购", "盒马 原味酸奶"]):
        return category(("买买买", "家庭购物"))
    if has_any(text, ["短袖", "长袖", "裤子", "裙子", "外套", "衬衫", "t恤", "体恤", "卫衣", "羽绒服", "唐狮", "优衣库", "desigual", "美洲城(澳門)戶外用品一人有限公司"]):
        return category(("买买买", "衣服裤子"))
    if has_any(text, ["cpb", "肌肤之钥", "妆前乳", "芙丽芳丝", "洁面泡沫", "口红", "粉底", "眼影", "睫毛膏"]):
        return category(("买买买", "化妆品"))
    if has_any(text, ["ubras", "乳贴", "果冻硅胶", "面膜", "精华", "防晒", "洗面奶", "水乳", "欧莱雅", "无印良品muji"]):
        return category(("买买买", "护肤品"))
    if "花" in text and has_any(text, ["💐", "花束", "鲜花"]):
        return category(("买买买", "花💐"))
    if has_any(text, ["代购"]):
        return category(("买买买", "代购"))
    if has_any(text, ["格力海岸"]):
        return category(("买买买", "🏠格力海岸"))
    if has_any(text, ["快乐的东西"]):
        return category(("买买买", "快乐的东西（2020算在玩得开心）"))
    if has_any(text, ["fluffy", "ringo", "kkv", "kk-store", "wow colour", "手账", "高光笔", "贴纸", "离型纸", "小红书", "中免日上"]):
        return category(("买买买", "其他"))

    return category(("其他杂项", "其他支出"))


def classify_income(row: dict) -> tuple[str, str]:
    text = normalized_text(row)
    amount = Decimal(str(row.get("金额") or "0"))

    if has_any(text, ["龙衔", "卢翮"]):
        return category(("智力活动", "副业收入"))
    if has_any(text, ["人才补贴", "补贴"]):
        return category(("智力活动", "人才补贴"))
    if has_any(text, ["奖金"]):
        return category(("智力活动", "奖金收入"))
    if has_any(text, ["工资"]):
        return category(("智力活动", "工资收入"))
    if has_any(text, ["公积金"]):
        return category(("智力活动", "公积金"))
    if has_any(text, ["红包"]) and amount < Decimal("1"):
        return category(("智力活动", "红包"))
    if has_any(text, ["aa", "aa制", "因因", "群收款"]):
        return category(("其他收入", "Aa制还钱"))
    if has_any(text, ["享瘦健康"]):
        return category(("其他收入", "理财收入"))
    if has_any(text, ["浅蓝色", "礼金"]):
        return category(("其他收入", "礼金收入"))
    if has_any(text, ["中奖", "福彩", "体彩"]):
        return category(("其他收入", "中奖收入"))
    if has_any(text, ["保险", "理赔", "航延"]):
        return category(("其他收入", "保险"))
    if has_any(text, ["房租"]):
        return category(("其他收入", "🏠房租"))
    if has_any(text, ["他人借钱买房"]):
        return category(("其他收入", "他人借钱买🏠"))
    return category(("其他收入", "其他收入"))


def map_account(row: dict) -> str:
    payment = str(row.get("收/付款方式") or "").strip()
    main = payment.split("&", 1)[0].strip()

    if not main or main == "/":
        account = "微信钱"
    elif "交通银行储蓄卡" in main:
        account = "交行卡"
    elif "交通银行信用卡" in main:
        account = "交行6.8还"
    elif "招商银行储蓄卡" in main:
        account = "招行"
    elif "招商银行信用卡(6020)" in main:
        cutoff = datetime(parse_dt(row["交易时间"]).year, 6, 15)
        account = "招行8.2还款" if parse_dt(row["交易时间"]) >= cutoff else "招行7.2还款"
    elif main == "余额宝":
        account = "余额宝"
    elif main == "零钱":
        account = "微信钱"
    elif main in {"支付宝", "账户余额"} or "支付宝小荷包" in main or main == "猫猫饭卡":
        account = "支付宝"
    else:
        account = "支付宝"

    if account not in VALID_ACCOUNTS:
        raise ValueError(f"账户不存在于分类体系: {account}（付款方式：{payment}）")
    return account


def map_investment_account(row: dict) -> str:
    text = normalized_text(row)
    if has_any(text, ["上海金", "黄金"]):
        account = GOLD_INVESTMENT_ACCOUNT
    elif has_any(text, ["纳斯达克", "纳指"]):
        account = NASDAQ_INVESTMENT_ACCOUNT
    else:
        raise ValueError(f"无法识别理财转入账户: {row.get('备注') or row.get('商品说明') or ''}")
    if account not in VALID_ACCOUNTS:
        raise ValueError(f"账户不存在于分类体系: {account}")
    return account


def build_suishouji_entry(row: dict) -> list:
    direction = row["收/支"]
    if direction == "支出":
        first, second = classify_expense(row)
    elif direction == "收入":
        first, second = classify_income(row)
    elif direction == "转账":
        return [
            "转账",
            str(row["交易时间"]),
            map_account(row),
            map_investment_account(row),
            float(row["金额"]),
            "",
            "",
            "",
            str(row.get("备注") or ""),
        ]
    else:
        raise ValueError(f"不支持的收/支类型: {direction}")
    return [
        direction,
        str(row["交易时间"]),
        first,
        second,
        map_account(row),
        float(row["金额"]),
        "",
        "",
        "",
        str(row.get("备注") or ""),
    ]


def read_merged_rows(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True)["合并账单"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def convert_with_soffice(input_path: Path, output_dir: Path, fmt: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.gettempdir()) / "suishouji_lo_profile"
    profile.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["XDG_CACHE_HOME"] = str(Path(tempfile.gettempdir()) / "fontconfig-cache")
    command = [
        resolve_soffice(),
        f"-env:UserInstallation=file://{profile}",
        "--headless",
        "--convert-to",
        fmt,
        "--outdir",
        str(output_dir),
        str(input_path),
    ]
    subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env=env)
    extension = fmt.split(":", 1)[0]
    converted = output_dir / f"{input_path.stem}.{extension}"
    if not converted.exists():
        raise FileNotFoundError(f"转换后文件不存在: {converted}")
    return converted


def ensure_xlsx_template(template_path: Path, work_dir: Path) -> Path:
    if template_path.suffix.lower() == ".xlsx":
        target = work_dir / template_path.name
        shutil.copy2(template_path, target)
        return target
    return convert_with_soffice(template_path, work_dir, "xlsx")


def copy_row_style(sheet, source_row: int, target_row: int, max_col: int = 10) -> None:
    for col in range(1, max_col + 1):
        source = sheet.cell(source_row, col)
        target = sheet.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
    if sheet.row_dimensions[source_row].height:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def clear_data_rows(sheet, max_col: int = 10) -> None:
    for row in range(2, sheet.max_row + 1):
        for col in range(1, max_col + 1):
            sheet.cell(row, col).value = None


def fill_sheet(sheet, entries: list[list]) -> None:
    clear_data_rows(sheet)
    for offset, entry in enumerate(entries, start=2):
        copy_row_style(sheet, 2, offset)
        for col, value in enumerate(entry, start=1):
            sheet.cell(offset, col).value = value


def write_report(rows: list[dict], output_path: Path, report_path: Path, start_time: str) -> None:
    expense = [row for row in rows if row["收/支"] == "支出"]
    income = [row for row in rows if row["收/支"] == "收入"]
    transfers = [row for row in rows if row["收/支"] == "转账"]
    transfer_destinations = Counter(build_suishouji_entry(row)[3] for row in transfers)
    transfer_summary = "，".join(
        f"{account} {count} 笔" for account, count in sorted(transfer_destinations.items())
    )
    supplement_reminders = manual_supplement_reminders(rows, start_time)
    lines = [
        "# 随手记录入检查 - 2026-07-06",
        "",
        f"- 起始时间：`{start_time}`，包含这条",
        f"- 输出文件：`{output_path}`",
        f"- 本次写入：{len(rows)} 笔",
        f"- 支出：{len(expense)} 笔",
        f"- 收入：{len(income)} 笔",
        f"- 转账：{len(transfers)} 笔（支付宝理财购买；{transfer_summary or '无'}）",
        "- 模板示例行：已删除，真实数据从第 2 行开始写入",
        f"- 账户检查提醒：{ACCOUNT_REVIEW_REMINDER}",
        "",
        "## 手动补记提醒",
        "",
    ]
    if supplement_reminders:
        lines.extend(f"- {reminder}" for reminder in supplement_reminders)
    else:
        lines.append("- 本次账单日期范围未覆盖工资、公积金或房租的常见发生日期。")
    lines.extend(
        [
            "",
            "## 逐笔录入明细",
            "",
        ]
    )
    for row in rows:
        entry = build_suishouji_entry(row)
        if row["收/支"] == "转账":
            lines.append(
                f"- {entry[1]}｜转账｜{entry[4]:.2f}｜{entry[2]} → {entry[3]}｜{row.get('交易对象') or ''}｜{row.get('商品说明') or ''}"
            )
        else:
            lines.append(
                f"- {entry[1]}｜{entry[0]}｜{entry[5]:.2f}｜{entry[4]}｜{entry[2]} / {entry[3]}｜{row.get('交易对象') or ''}｜{row.get('商品说明') or ''}"
            )
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def fill_suishouji_template(
    bills_path: Path,
    template_path: Path,
    output_path: Path,
    report_path: Path,
    start_time: str,
) -> tuple[list[dict], Path]:
    rows = filter_rows_from_start(read_merged_rows(bills_path), start_time)
    expense_entries = [build_suishouji_entry(row) for row in rows if row["收/支"] == "支出"]
    income_entries = [build_suishouji_entry(row) for row in rows if row["收/支"] == "收入"]
    transfer_entries = [build_suishouji_entry(row) for row in rows if row["收/支"] == "转账"]

    with tempfile.TemporaryDirectory(prefix="suishouji_fill_") as tmp:
        work_dir = Path(tmp)
        template_xlsx = ensure_xlsx_template(template_path, work_dir)
        workbook = load_workbook(template_xlsx)
        fill_sheet(workbook["支出"], expense_entries)
        fill_sheet(workbook["收入"], income_entries)
        fill_sheet(workbook["转账"], transfer_entries)
        working_xlsx = work_dir / f"{output_path.stem}.xlsx"
        workbook.save(working_xlsx)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        if output_path.suffix.lower() == ".xlsx":
            shutil.copy2(working_xlsx, output_path)
        else:
            converted = convert_with_soffice(working_xlsx, work_dir, output_path.suffix.lower().lstrip("."))
            shutil.copy2(converted, output_path)

    write_report(rows, output_path, report_path, start_time)
    return rows, output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="将合并账单分类后写入随手记模板。")
    parser.add_argument("--bills", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--start", required=True)
    args = parser.parse_args()

    rows, output = fill_suishouji_template(args.bills, args.template, args.output, args.report, args.start)
    expense_count = sum(1 for row in rows if row["收/支"] == "支出")
    income_count = sum(1 for row in rows if row["收/支"] == "收入")
    transfer_count = sum(1 for row in rows if row["收/支"] == "转账")
    print(f"写入 {len(rows)} 笔：支出 {expense_count}，收入 {income_count}，转账 {transfer_count}")
    print(f"输出文件: {output}")
    print(f"检查记录: {args.report}")
    print(ACCOUNT_REVIEW_REMINDER)
    reminders = manual_supplement_reminders(rows, args.start)
    if reminders:
        print("手动补记提醒:")
        for reminder in reminders:
            print(f"- {reminder}")


if __name__ == "__main__":
    main()
