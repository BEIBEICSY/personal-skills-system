from __future__ import annotations

import argparse
import csv
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_COLUMNS = [
    "来源",
    "源文件行",
    "交易时间",
    "交易类型",
    "交易对象",
    "商品说明",
    "收/支",
    "金额",
    "收/付款方式",
    "当前状态",
    "交易订单号",
    "商家订单号",
    "原始备注",
    "备注",
]


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\t", "").strip()


def parse_amount(value) -> Decimal:
    text = clean_text(value).replace(",", "").replace("¥", "")
    if not text:
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"金额无法解析: {value!r}") from exc


def normalize_time(value) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return text


def build_note(counterparty: str, product: str, raw_note: str) -> str:
    parts = [clean_text(counterparty), clean_text(product), clean_text(raw_note)]
    normalized = [p for p in parts if p and p != "/"]
    return " | ".join(normalized)


def is_supported_alipay_investment_purchase(transaction_type: str, product: str) -> bool:
    return (
        clean_text(transaction_type) == "投资理财"
        and "买入" in clean_text(product)
        and any(keyword in clean_text(product) for keyword in ["纳斯达克", "纳指", "上海金", "黄金"])
    )


def normalize_alipay_row(row: dict, source_row: int) -> dict:
    counterparty = clean_text(row.get("交易对方"))
    product = clean_text(row.get("商品说明"))
    raw_note = clean_text(row.get("备注"))
    transaction_type = clean_text(row.get("交易分类"))
    direction = clean_text(row.get("收/支"))
    if is_supported_alipay_investment_purchase(transaction_type, product):
        direction = "转账"
    return {
        "来源": "支付宝",
        "源文件行": source_row,
        "交易时间": normalize_time(row.get("交易时间")),
        "交易类型": transaction_type,
        "交易对象": counterparty,
        "商品说明": product,
        "收/支": direction,
        "金额": parse_amount(row.get("金额")),
        "收/付款方式": clean_text(row.get("收/付款方式")),
        "当前状态": clean_text(row.get("交易状态")),
        "交易订单号": clean_text(row.get("交易订单号")),
        "商家订单号": clean_text(row.get("商家订单号")),
        "原始备注": raw_note,
        "备注": build_note(counterparty, product, raw_note),
    }


def normalize_wechat_row(row: dict, source_row: int) -> dict:
    counterparty = clean_text(row.get("交易对方"))
    product = clean_text(row.get("商品"))
    raw_note = clean_text(row.get("备注"))
    return {
        "来源": "微信",
        "源文件行": source_row,
        "交易时间": normalize_time(row.get("交易时间")),
        "交易类型": clean_text(row.get("交易类型")),
        "交易对象": counterparty,
        "商品说明": product,
        "收/支": clean_text(row.get("收/支")),
        "金额": parse_amount(row.get("金额(元)")),
        "收/付款方式": clean_text(row.get("支付方式")),
        "当前状态": clean_text(row.get("当前状态")),
        "交易订单号": clean_text(row.get("交易单号")),
        "商家订单号": clean_text(row.get("商户单号")),
        "原始备注": raw_note,
        "备注": build_note(counterparty, product, raw_note),
    }


def read_alipay(path: Path) -> list[dict]:
    text = path.read_bytes().decode("gb18030")
    lines = text.splitlines()
    header_index = next(i for i, line in enumerate(lines) if line.startswith("交易时间,"))
    rows = []
    for offset, row in enumerate(csv.DictReader(lines[header_index:]), start=header_index + 2):
        if clean_text(row.get("交易时间")):
            rows.append(normalize_alipay_row(row, offset))
    return rows


def read_wechat(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    header_row = 18
    headers = [cell.value for cell in sheet[header_row]]
    rows = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not values or not values[0]:
            continue
        rows.append(normalize_wechat_row(dict(zip(headers, values)), source_row))
    return rows


def contains_refund_text(record: dict) -> bool:
    text = "".join(
        [
            record.get("交易类型", ""),
            record.get("交易对象", ""),
            record.get("商品说明", ""),
            record.get("当前状态", ""),
            record.get("原始备注", ""),
        ]
    )
    return "退款" in text


def base_refund_order(order: str) -> str:
    text = clean_text(order)
    if "_" in text:
        text = text.split("_", 1)[0]
    if "*" in text:
        text = text.split("*", 1)[0]
    return text


def initial_exclusion_reason(record: dict) -> str:
    status = record.get("当前状态", "")
    direction = record.get("收/支", "")
    if status in {"退款成功"} or contains_refund_text(record):
        if direction == "支出" and status in {"已全额退款"}:
            return "退款抵消原支出"
        return "退款相关"
    if status in {"已全额退款"}:
        return "退款抵消原支出" if direction == "支出" else "退款相关"
    if direction not in {"支出", "收入", "转账"}:
        return "不计收支/中性交易"
    if direction == "转账" and not is_supported_alipay_investment_purchase(
        record.get("交易类型", ""), record.get("商品说明", "")
    ):
        return "不计收支/中性交易"
    if status in {"交易关闭"}:
        return "交易关闭"
    return ""


def is_same_money(left: Decimal, right: Decimal) -> bool:
    return left.quantize(Decimal("0.01")) == right.quantize(Decimal("0.01"))


def refund_matches_expense(refund: dict, expense: dict) -> bool:
    if refund.get("来源") != expense.get("来源") or expense.get("收/支") != "支出":
        return False
    if not is_same_money(refund["金额"], expense["金额"]):
        return False

    refund_order = clean_text(refund.get("交易订单号"))
    expense_order = clean_text(expense.get("交易订单号"))
    refund_merchant = clean_text(refund.get("商家订单号"))
    expense_merchant = clean_text(expense.get("商家订单号"))
    refund_base = base_refund_order(refund_order)

    order_match = bool(refund_base and expense_order and refund_base == expense_order)
    merchant_match = bool(refund_merchant and expense_merchant and refund_merchant == expense_merchant)
    wechat_redpacket_match = bool(refund_order and expense_merchant and refund_order == expense_merchant)
    counterparty_match = bool(refund.get("交易对象") and refund.get("交易对象") == expense.get("交易对象"))

    if order_match or merchant_match or wechat_redpacket_match:
        return True
    return counterparty_match and refund.get("商品说明") == expense.get("商品说明")


def filter_clean_records(records: Iterable[dict]) -> tuple[list[dict], list[dict]]:
    rows = [dict(record) for record in records]
    reasons = {id(record): initial_exclusion_reason(record) for record in rows}
    refund_rows = [record for record in rows if reasons[id(record)] == "退款相关"]

    for refund in refund_rows:
        for expense in rows:
            if refund is expense:
                continue
            if refund_matches_expense(refund, expense):
                reasons[id(expense)] = "退款抵消原支出"

    kept = []
    excluded = []
    for record in rows:
        reason = reasons[id(record)]
        if reason:
            item = dict(record)
            item["排除原因"] = reason
            excluded.append(item)
        else:
            kept.append(record)

    kept.sort(key=lambda row: row["交易时间"], reverse=True)
    excluded.sort(key=lambda row: row["交易时间"], reverse=True)
    return kept, excluded


def write_merged_workbook(rows: list[dict], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合并账单"
    sheet.append(OUTPUT_COLUMNS)

    for record in rows:
        sheet.append(
            [
                record["来源"],
                record["源文件行"],
                record["交易时间"],
                record["交易类型"],
                record["交易对象"],
                record["商品说明"],
                record["收/支"],
                float(record["金额"]),
                record["收/付款方式"],
                record["当前状态"],
                record["交易订单号"],
                record["商家订单号"],
                record["原始备注"],
                record["备注"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 10,
        "B": 10,
        "C": 20,
        "D": 16,
        "E": 24,
        "F": 42,
        "G": 10,
        "H": 12,
        "I": 24,
        "J": 16,
        "K": 30,
        "L": 30,
        "M": 24,
        "N": 60,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet["H"][1:]:
        cell.number_format = "0.00"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def format_money(value: Decimal) -> str:
    return f"{value.quantize(Decimal('0.01'))}"


def record_line(record: dict) -> str:
    return (
        f"- {record['交易时间']}｜{record['来源']}｜{record['收/支']}｜"
        f"{format_money(record['金额'])}｜{record['交易对象']}｜{record['商品说明']}｜"
        f"{record.get('排除原因', '录入')}"
    )


def write_check_report(kept: list[dict], excluded: list[dict], output_path: Path, workbook_path: Path) -> None:
    kept_counter = Counter((row["来源"], row["收/支"]) for row in kept)
    excluded_counter = Counter(row["排除原因"] for row in excluded)
    refund_excluded = [row for row in excluded if "退款" in row["排除原因"]]

    lines = [
        "# 合并账单清洗记录 - 2026-07-06",
        "",
        f"- 合并账单输出：`{workbook_path}`",
        f"- 录入合并账单：{len(kept)} 笔",
        f"- 排除记录：{len(excluded)} 笔",
        "",
        "## 录入统计",
        "",
    ]
    for key, count in sorted(kept_counter.items()):
        lines.append(f"- {key[0]} {key[1]}：{count} 笔")

    lines.extend(["", "## 排除原因统计", ""])
    for reason, count in sorted(excluded_counter.items()):
        lines.append(f"- {reason}：{count} 笔")

    lines.extend(["", "## 退款抵消和退款排除明细", ""])
    lines.extend(record_line(row) for row in refund_excluded)

    lines.extend(["", "## 已录入明细", ""])
    lines.extend(record_line(row) for row in kept)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def merge_bills(alipay_path: Path, wechat_path: Path, output_path: Path, report_path: Path) -> tuple[list[dict], list[dict]]:
    records = read_alipay(alipay_path) + read_wechat(wechat_path)
    kept, excluded = filter_clean_records(records)
    write_merged_workbook(kept, output_path)
    write_check_report(kept, excluded, report_path, output_path)
    return kept, excluded


def main() -> None:
    parser = argparse.ArgumentParser(description="合并并清洗支付宝、微信账单。")
    parser.add_argument("--alipay", required=True, type=Path)
    parser.add_argument("--wechat", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()

    kept, excluded = merge_bills(args.alipay, args.wechat, args.output, args.report)
    print(f"录入 {len(kept)} 笔，排除 {len(excluded)} 笔")
    print(f"合并账单: {args.output}")
    print(f"检查记录: {args.report}")


if __name__ == "__main__":
    main()
